#!/usr/bin/env python3
"""Create Nobl9 SLO YAML from an Excel creation template.

Dependencies: python -m pip install openpyxl PyYAML
Example:
  python slo_sheet_to_yaml.py examples/sample_slos.xlsx --output-dir generated
  python slo_sheet_to_yaml.py examples/sample_slos.xlsx --output-dir generated \
    --compare-original baseline-latency.yaml baseline-reliability.yaml \
    --report comparison.txt

One row = one objective. Group by (project, slo_name). This version supports
Splunk Observability, Occurrences, and one rolling window. Latency raw programs
are generated, never copied from raw_program. Reliability good/total programs
are read from the sheet. Context/query disagreements are validation errors.
No API calls are made, no SLOs are applied, and existing output files require
--overwrite. Configuration comparison is semantic, not byte-for-byte.
"""

import argparse
import copy
import difflib
import json
import math
import re
import sys
from pathlib import Path

import openpyxl
import yaml


class SourceLoader(yaml.SafeLoader):
    pass


# Nobl9 exports JSON-style scientific numbers such as 1e+09. PyYAML's
# default YAML 1.1 resolver otherwise treats these as strings.
SourceLoader.add_implicit_resolver(
    'tag:yaml.org,2002:float',
    re.compile(r'^[-+]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)[eE][-+]?[0-9]+$'),
    list('-+0123456789.'),
)


SHARED = (
    'slo_name slo_display_name project service description data_source_name '
    'data_source_project data_source_kind budgeting_method window_unit '
    'window_count window_is_rolling'
).split()
OBJECTIVE = (
    'endpoint http_method objective_name objective_display_name target '
    'objective_value metric_type operator incremental splunk_environment '
    'splunk_service splunk_operation raw_program good_program total_program'
).split()
REQUIRED_HEADERS = {
    'endpoint', 'http_method', 'slo_name', 'project', 'service',
    'data_source_name', 'data_source_kind', 'budgeting_method', 'window_unit',
    'window_count', 'window_is_rolling', 'objective_name', 'target',
    'objective_value', 'metric_type', 'operator', 'incremental',
    'splunk_environment', 'splunk_service', 'splunk_operation',
    'good_program', 'total_program',
}


class ValidationError(ValueError):
    pass


def text(value):
    return '' if value is None else str(value).strip()


def number(value, field, row):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValidationError(f'Row {row}: {field} must be a numeric Excel cell.')
    if not math.isfinite(value):
        raise ValidationError(f'Row {row}: {field} must be finite.')
    return value


def boolean(value, field, row):
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in ('true', 'false'):
        return value.strip().lower() == 'true'
    raise ValidationError(f'Row {row}: {field} must be TRUE or FALSE.')


def name(value, field, row):
    value = text(value)
    if len(value) > 63 or not re.fullmatch(r'[a-z0-9](?:[a-z0-9-]*[a-z0-9])?', value):
        raise ValidationError(f'Row {row}: {field} must be a DNS-style name of 1–63 characters.')
    return value


def display(value, field, row):
    value = text(value)
    if len(value) > 63:
        raise ValidationError(f'Row {row}: {field} exceeds 63 characters.')
    return value


def quoted(value):
    # Escape values as SignalFlow Python-compatible string literals.
    return repr(value)


def raw_program(row, percentile):
    f = (
        f"filter('sf_environment', {quoted(row['splunk_environment'])}) and "
        f"filter('sf_service', {quoted(row['splunk_service'])}) and "
        f"filter('sf_operation', {quoted(row['splunk_operation'])}) and "
        "filter('sf_kind', 'SERVER', 'CONSUMER') and "
        "filter('sf_error', 'false') and (not filter('sf_dimensionalized', '*')) "
        "and (not filter('sf_serviceMesh', '*'))"
    )
    return (
        f'filter_ = {f}\n\n'
        f"A = histogram('spans', filter=filter_).percentile(pct={percentile}).publish(label='p{percentile}')"
    )


def normalized_program(value):
    return '\n'.join(line.rstrip() for line in str(value).replace('\r\n', '\n').split('\n')).strip()


def filter_values(program, key):
    # Inspect static filter string literals without executing spreadsheet code.
    import ast
    try:
        tree = ast.parse(program)
    except SyntaxError as exc:
        raise ValidationError(f'Invalid SignalFlow program syntax: {exc.msg}') from exc
    return [
        node.args[1].value for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name) and node.func.id == 'filter'
        and len(node.args) >= 2
        and isinstance(node.args[0], ast.Constant) and node.args[0].value == key
        and isinstance(node.args[1], ast.Constant) and isinstance(node.args[1].value, str)
    ]


def check_context(program, row, label):
    for key, column in [('sf_environment', 'splunk_environment'),
                        ('sf_service', 'splunk_service'), ('sf_operation', 'splunk_operation')]:
        values = filter_values(program, key)
        if not values or any(v != row[column] for v in values):
            raise ValidationError(f"Row {row['_row']}: {label} {key} disagrees with {column}.")


def read_rows(path, sheet):
    book = openpyxl.load_workbook(path, data_only=False)
    if sheet not in book.sheetnames:
        raise ValidationError(f'Sheet {sheet!r} does not exist.')
    ws = book[sheet]
    headers = [text(c.value) for c in ws[1]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValidationError('Headers must be nonblank and unique.')
    missing = REQUIRED_HEADERS - set(headers)
    unknown = set(headers) - set(SHARED + OBJECTIVE) - {h for h in headers if h.startswith('label_')}
    if missing or unknown:
        raise ValidationError(f'Invalid headers; missing={sorted(missing)}, unknown={sorted(unknown)}')
    rows = []
    for cells in ws.iter_rows(min_row=2):
        if all(c.value is None or c.value == '' for c in cells):
            continue
        for c in cells:
            if c.data_type == 'f':
                raise ValidationError(f'Cell {c.coordinate}: formulas are not supported; provide literal values.')
        r = dict(zip(headers, (c.value for c in cells)))
        r['_row'] = cells[0].row
        for field in SHARED + OBJECTIVE:
            r.setdefault(field, None)
        rows.append(r)
    if not rows:
        raise ValidationError('No objective rows found.')
    return rows, [h for h in headers if h.startswith('label_')]


def build_slos(rows, label_columns, percentile=95, allow_reference_differences=False):
    groups = {}; messages = []
    for r in rows:
        n = r['_row']
        for field in ['slo_name', 'project', 'service', 'data_source_name', 'objective_name']:
            r[field] = name(r[field], field, n)
        if text(r['data_source_project']):
            r['data_source_project'] = name(r['data_source_project'], 'data_source_project', n)
        else:
            r['data_source_project'] = r['project']
        for field in ['slo_display_name', 'objective_display_name']:
            r[field] = display(r[field], field, n)
        for field in ['endpoint', 'http_method', 'description', 'data_source_kind',
                      'budgeting_method', 'window_unit', 'metric_type', 'operator',
                      'splunk_environment', 'splunk_service', 'splunk_operation']:
            r[field] = text(r[field])
        if len(r['description']) > 1050:
            raise ValidationError(f'Row {n}: description exceeds 1050 characters.')
        if r['data_source_kind'] not in ('Direct', 'Agent'):
            raise ValidationError(f'Row {n}: data_source_kind must be Direct or Agent.')
        if r['budgeting_method'] != 'Occurrences':
            raise ValidationError(f'Row {n}: only Occurrences budgeting is supported.')
        if r['window_unit'] not in ('Minute', 'Hour', 'Day'):
            raise ValidationError(f'Row {n}: invalid rolling window unit.')
        count = number(r['window_count'], 'window_count', n)
        if count <= 0 or int(count) != count:
            raise ValidationError(f'Row {n}: window_count must be a positive integer.')
        r['window_count'] = int(count)
        r['window_is_rolling'] = boolean(r['window_is_rolling'], 'window_is_rolling', n)
        if not r['window_is_rolling']:
            raise ValidationError(f'Row {n}: only rolling windows are supported.')
        if not r['endpoint'].startswith('/') or r['http_method'] not in {
                'GET', 'POST', 'PUT', 'PATCH', 'DELETE', 'HEAD', 'OPTIONS'}:
            raise ValidationError(f'Row {n}: invalid endpoint or HTTP method.')
        if r['splunk_operation'] != r['http_method'] + ' ' + r['endpoint']:
            raise ValidationError(f'Row {n}: splunk_operation must equal http_method + space + endpoint.')
        if not r['splunk_environment'] or not r['splunk_service']:
            raise ValidationError(f'Row {n}: Splunk environment and service are required.')
        target = number(r['target'], 'target', n)
        if not 0 <= target < 1:
            raise ValidationError(f'Row {n}: target must be a fraction in [0, 1), e.g. 0.999.')
        value = number(r['objective_value'], 'objective_value', n)
        labels = {}
        for col in label_columns:
            key = col[6:]
            if len(key) > 63 or not re.fullmatch(r'[a-z](?:[a-z0-9_-]*[a-z0-9])?', key):
                raise ValidationError(f'Invalid label column {col!r}.')
            if text(r.get(col)):
                values = [v.strip() for v in text(r[col]).split('|')]
                if any(not v or len(v) > 200 for v in values):
                    raise ValidationError(f'Row {n}: invalid values in {col}.')
                labels[key] = values
        metadata = {'name': r['slo_name'], 'project': r['project']}
        if r['slo_display_name']:
            metadata['displayName'] = r['slo_display_name']
        if labels:
            metadata['labels'] = labels
        source = {'name': r['data_source_name'], 'project': r['data_source_project'], 'kind': r['data_source_kind']}
        spec = {'service': r['service'], 'indicator': {'metricSource': source},
                'budgetingMethod': r['budgeting_method'],
                'timeWindows': [{'unit': r['window_unit'], 'count': r['window_count'], 'isRolling': True}]}
        if r['description']:
            spec['description'] = r['description']
        shared = {'apiVersion': 'n9/v1alpha', 'kind': 'SLO', 'metadata': metadata, 'spec': spec}
        key = (r['project'], r['slo_name'])
        if key not in groups:
            groups[key] = {'shared': shared, 'rows': [], 'objectives': []}
        group = groups[key]
        if shared != group['shared']:
            raise ValidationError(f'Row {n}: shared SLO fields conflict with row {group["rows"][0]} for {key}.')
        objective = {'name': r['objective_name'], 'target': target, 'value': value}
        if r['objective_display_name']:
            objective['displayName'] = r['objective_display_name']
        if r['metric_type'] == 'rawMetric':
            if r['operator'] not in ('lt', 'lte', 'gt', 'gte'):
                raise ValidationError(f'Row {n}: rawMetric requires a valid operator.')
            if text(r['good_program']) or text(r['total_program']) or r['incremental'] not in (None, ''):
                raise ValidationError(f'Row {n}: countMetrics fields must be blank for rawMetric.')
            program = raw_program(r, percentile)
            if text(r['raw_program']):
                match = normalized_program(r['raw_program']) == normalized_program(program)
                messages.append(f'Row {n}: generated raw program vs spreadsheet reference: {"MATCH" if match else "DIFFERENT"}')
                if not match and not allow_reference_differences:
                    raise ValidationError(f'Row {n}: generated raw program differs from reference. Review it; use --allow-reference-differences for intentional changes.')
            else:
                messages.append(f'Row {n}: generated raw program; no spreadsheet reference supplied.')
            objective.update(op=r['operator'], rawMetric={'query': {'splunkObservability': {'program': program}}})
        elif r['metric_type'] == 'countMetrics':
            if r['operator'] or text(r['raw_program']):
                raise ValidationError(f'Row {n}: rawMetric fields must be blank for countMetrics.')
            counts = {'incremental': boolean(r['incremental'], 'incremental', n)}
            for field, branch in [('good_program', 'good'), ('total_program', 'total')]:
                if not text(r[field]):
                    raise ValidationError(f'Row {n}: {field} is required.')
                program = str(r[field]).replace('\r\n', '\n')
                check_context(program, r, field)
                counts[branch] = {'splunkObservability': {'program': program}}
            objective['countMetrics'] = counts
        else:
            raise ValidationError(f'Row {n}: metric_type must be rawMetric or countMetrics.')
        if any(o['name'] == objective['name'] for o in group['objectives']):
            raise ValidationError(f'Row {n}: duplicate objective name within {key}.')
        if any(o['value'] == objective['value'] for o in group['objectives']):
            raise ValidationError(f'Row {n}: duplicate objective value within {key}.')
        group['rows'].append(n); group['objectives'].append(objective)
        if len(group['objectives']) > 12:
            raise ValidationError(f'Row {n}: more than 12 objectives in {key}.')
    slos = []
    for g in groups.values():
        slo = copy.deepcopy(g['shared']); slo['spec']['objectives'] = g['objectives']; slos.append(slo)
    return slos, messages


def creation_config(slo):
    """Remove only known export-only fields; retain everything else for comparison."""
    result = copy.deepcopy(slo)
    result.pop('status', None)
    spec = result.get('spec', {})
    spec.pop('createdAt', None); spec.pop('createdBy', None)
    if spec.get('description') == '':
        spec.pop('description')
    for window in spec.get('timeWindows', []):
        if window.get('isRolling') is True:
            window.pop('period', None)
    return result


def compare_originals(slos, files):
    originals = {}
    for path in files:
        for doc in yaml.load_all(Path(path).read_text(encoding='utf-8'), Loader=SourceLoader):
            if doc is None:
                continue
            for slo in doc if isinstance(doc, list) else [doc]:
                if not isinstance(slo, dict) or slo.get('kind') != 'SLO':
                    raise ValidationError(f'{path}: expected SLO objects.')
                key = (slo['metadata']['project'], slo['metadata']['name'])
                if key in originals:
                    raise ValidationError(f'Duplicate original SLO: {key}')
                originals[key] = creation_config(slo)
    generated = {(s['metadata']['project'], s['metadata']['name']): creation_config(s) for s in slos}
    lines = [
        'Creation configuration comparison (parsed YAML, not text formatting)',
        'Excluded: status; spec.createdAt; spec.createdBy; rolling timeWindows[].period.',
        'An empty optional spec.description is treated as omitted.',
        'All remaining fields and objective ordering are compared.',
    ]
    matches = True
    for key in sorted(set(originals) | set(generated)):
        if key not in originals or key not in generated:
            lines.append(f'MISMATCH {key}: missing from {"originals" if key not in originals else "generated YAML"}')
            matches = False
        elif originals[key] == generated[key]:
            lines.append(f'MATCH {key[0]}/{key[1]}: {len(generated[key]["spec"]["objectives"])} objectives')
        else:
            matches = False; lines.append(f'MISMATCH {key[0]}/{key[1]}')
            a = json.dumps(originals[key], indent=2, sort_keys=True, ensure_ascii=False).splitlines()
            b = json.dumps(generated[key], indent=2, sort_keys=True, ensure_ascii=False).splitlines()
            lines.extend(difflib.unified_diff(a, b, fromfile='original creation config', tofile='generated creation config', lineterm=''))
    lines.append('PASS: all creation configurations match.' if matches else 'FAIL: creation configurations differ.')
    return matches, lines


class Dumper(yaml.SafeDumper):
    def ignore_aliases(self, data):
        return True


def represent_string(dumper, value):
    return dumper.represent_scalar('tag:yaml.org,2002:str', value, style='|' if '\n' in value else None)


Dumper.add_representer(str, represent_string)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('spreadsheet', type=Path)
    parser.add_argument('--sheet', default='SLO Template')
    parser.add_argument('--output-dir', type=Path, default=Path('generated'))
    parser.add_argument('--latency-percentile', type=int, default=95,
                        help='Percentile used for generated raw programs (default: 95).')
    parser.add_argument('--allow-reference-differences', action='store_true',
                        help='Allow generated raw programs to differ from old spreadsheet references.')
    parser.add_argument('--compare-original', nargs='+', type=Path, default=[])
    parser.add_argument('--report', type=Path)
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args(argv)
    try:
        if not 0 < args.latency_percentile <= 100:
            raise ValidationError('--latency-percentile must be 1–100.')
        rows, labels = read_rows(args.spreadsheet, args.sheet)
        slos, lines = build_slos(rows, labels, args.latency_percentile, args.allow_reference_differences)
        payloads = []
        for slo in slos:
            # Include project in filename so identical SLO names in different projects cannot collide.
            filename = f"{slo['metadata']['project']}__{slo['metadata']['name']}.yaml"
            path = args.output_dir / filename
            content = yaml.dump([slo], Dumper=Dumper, sort_keys=False, allow_unicode=True, width=120)
            parsed = yaml.safe_load(content)
            if parsed != [slo]:
                raise ValidationError(f'Serialized YAML differs from in-memory configuration: {path}')
            payloads.append((path, content, parsed[0]))
        report_lines = ['Spreadsheet-to-YAML validation', f'{len(rows)} objective rows -> {len(slos)} SLOs.', *lines]
        match = True
        if args.compare_original:
            # Compare the parsed output bytes, not just the builder's in-memory objects.
            match, comparison = compare_originals([p[2] for p in payloads], args.compare_original)
            report_lines.extend(['', *comparison])
        report_lines.extend(['', 'Offline configuration validation only; no Nobl9 deployment or API reference check was performed.'])
        destinations = [p[0] for p in payloads] + ([args.report] if args.report else [])
        if len({p.resolve() for p in destinations}) != len(destinations):
            raise ValidationError('Output/report paths collide.')
        inputs = {args.spreadsheet.resolve(), *(p.resolve() for p in args.compare_original)}
        if any(p.resolve() in inputs for p in destinations):
            raise ValidationError('Output paths must not overwrite input files.')
        for path in destinations:
            if path.exists() and not args.overwrite:
                raise ValidationError(f'Output already exists: {path}; use --overwrite to replace generated files.')
        for path, content, _ in payloads:
            path.parent.mkdir(parents=True, exist_ok=True); path.write_text(content, encoding='utf-8')
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text('\n'.join(report_lines) + '\n', encoding='utf-8')
        print('\n'.join(report_lines))
        for path, _, _ in payloads:
            print(f'Created: {path}')
        return 0 if match else 1
    except (ValidationError, OSError, yaml.YAMLError, KeyError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
